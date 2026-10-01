import io

from docx import Document

import evidence_search
from evidence_search import (
    build_evidence_catalogue_docx,
    build_evidence_packets,
    group_evidence_by_source,
    quote_occurs_in_chunk,
    search_evidence,
    select_grounded_excerpt,
)


def _chunk(chunk_id="c1", source="India.docx", text=None, page=2, score=0.82):
    return {
        "chunk_id": chunk_id,
        "source_file": source,
        "text": text or (
            "GovRisk delivered an inter-agency forum in New Delhi. "
            "The programme strengthened financial investigations and asset recovery."
        ),
        "page_number": page,
        "relevance_score": score,
    }


def test_excerpt_is_exact_and_prefers_query_terms():
    text = (
        "The opening sentence concerns administration. "
        "The India programme strengthened financial investigations and asset recovery. "
        "The closing sentence concerns delivery."
    )
    excerpt = select_grounded_excerpt(text, "India financial investigations")
    assert "India programme" in excerpt
    assert quote_occurs_in_chunk(excerpt, text)


def test_quote_verifier_normalises_whitespace_but_rejects_paraphrase():
    chunk = "Evidence from India\ncovered financial   investigations."
    assert quote_occurs_in_chunk("Evidence from India covered financial investigations.", chunk)
    assert not quote_occurs_in_chunk("India delivered successful investigations.", chunk)


def test_packets_disclose_docx_locator_is_not_verified_page():
    packets = build_evidence_packets([_chunk()], "India investigations")
    assert len(packets) == 1
    packet = packets[0]
    assert packet["verification_status"] == "verified"
    assert packet["locator_type"] == "approximate_docx_segment"
    assert packet["locator_label"] == "DOCX extraction segment 2 (not a verified page)"


def test_pdf_page_is_labelled_as_pdf_page():
    packets = build_evidence_packets(
        [_chunk(source="General capability.pdf", page=7)],
        "financial investigations",
    )
    assert packets[0]["locator_type"] == "verified_pdf_page"
    assert packets[0]["locator_label"] == "PDF page 7"


def test_packets_limit_per_document_and_deduplicate_chunk_ids():
    chunks = [
        _chunk("c1", score=0.9),
        _chunk("c1", score=0.8),
        _chunk("c2", score=0.7),
        _chunk("c3", score=0.6),
    ]
    packets = build_evidence_packets(
        chunks,
        "India",
        excerpts_per_document=2,
    )
    assert [packet["chunk_id"] for packet in packets] == ["c1", "c2"]


def test_groups_preserve_retrieval_order_and_best_score():
    packets = build_evidence_packets(
        [
            _chunk("c1", "B.docx", score=0.5),
            _chunk("c2", "A.pdf", page=3, score=0.7),
            _chunk("c3", "B.docx", score=0.9),
        ],
        "India",
    )
    groups = group_evidence_by_source(packets)
    assert [group["source_name"] for group in groups] == ["B.docx", "A.pdf"]
    assert groups[0]["best_relevance_score"] == 0.9
    assert len(groups[0]["evidence"]) == 2


def test_search_uses_free_form_query_only_for_retrieval(monkeypatch):
    captured = {}

    def fake_retrieve(query, filters, top_k):
        captured.update(query=query, filters=filters, top_k=top_k)
        return {
            "retrieved_chunks": [_chunk()],
            "library_unavailable": False,
        }

    monkeypatch.setattr(evidence_search, "retrieve_query_chunks", fake_retrieve)
    result = search_evidence("  India and fraud  ", top_k=17)

    assert captured == {"query": "India and fraud", "filters": {}, "top_k": 17}
    assert result["documents_found"] == 1
    assert result["verified_evidence"] == 1
    assert result["needs_review"] == 0


def test_blank_request_never_calls_retrieval(monkeypatch):
    def fail(*args, **kwargs):
        raise AssertionError("retrieval must not run")

    monkeypatch.setattr(evidence_search, "retrieve_query_chunks", fail)
    result = search_evidence("  ")
    assert result["validation_error"] == "Enter an evidence request before searching."
    assert result["groups"] == []


def test_infrastructure_failure_is_not_reported_as_no_evidence(monkeypatch):
    monkeypatch.setattr(
        evidence_search,
        "retrieve_query_chunks",
        lambda *args, **kwargs: {
            "retrieved_chunks": [],
            "library_unavailable": True,
            "library_error": "ChromaUnavailableError: test",
        },
    )
    result = search_evidence("India")
    assert result["library_unavailable"] is True
    assert result["documents_found"] == 0
    assert "library_error" in result


def test_catalogue_contains_query_sources_quotes_and_chunk_ids():
    packets = build_evidence_packets(
        [
            _chunk("india-1", "India.docx"),
            _chunk("panama-1", "Panama.pdf", "The study addressed narcotics and organised crime.", 7),
        ],
        "India narcotics organised crime",
    )
    result = {
        "query": "India narcotics organised crime",
        "documents_found": 2,
        "verified_evidence": 2,
        "needs_review": 0,
        "groups": group_evidence_by_source(packets),
    }
    payload = build_evidence_catalogue_docx(result)
    document = Document(io.BytesIO(payload))
    text = "\n".join(p.text for p in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )

    assert payload.startswith(b"PK")
    assert "India narcotics organised crime" in text
    assert "India.docx" in text
    assert "Panama.pdf" in text
    assert "india-1" in text
    assert "panama-1" in text
    assert "not a verified page" in text
    assert "PDF page 7" in text


def test_catalogue_does_not_invent_project_metadata():
    packet = build_evidence_packets([_chunk()], "India")[0]
    result = {
        "query": "India",
        "documents_found": 1,
        "verified_evidence": 1,
        "needs_review": 0,
        "groups": group_evidence_by_source([packet]),
    }
    document = Document(io.BytesIO(build_evidence_catalogue_docx(result)))
    text = "\n".join(p.text for p in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "Unknown client" not in text
    assert "Unknown date" not in text

